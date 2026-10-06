"""Разбор выгрузок АИС: CSV и Excel, кодировка, разделитель, заголовки с дублями, приведение типов."""

from __future__ import annotations

import csv
import io
import itertools
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

from .field_maps import EntityMap, FieldSpec, norm_header

ENCODINGS = ("utf-8-sig", "cp1251")
DATE_FORMATS = ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%d.%m.%y")
TRUE_VALUES = {"1", "true", "t", "y", "yes", "да", "д"}
PAYMENT_TYPES = {"файл": "file", "ввод вручную": "manual", "зачисление из зарплаты": "salary",
                 "file": "file", "manual": "manual", "salary": "salary"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xls"}
# «1,2E+11» — так Excel записывает длинный номер, если ячейка была числом. Цифры уже потеряны.
_SCIENCE = re.compile(r"^[+-]?\d+(?:[.,]\d+)?[eE][+-]?\d+$")


class ConversionError(ValueError):
    """Ошибка приведения значения колонки."""


def decode(content: bytes) -> tuple[str, str]:
    for encoding in ENCODINGS:
        try:
            return content.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise ConversionError("Не удалось определить кодировку (ожидается UTF-8 или Windows-1251)")


def detect_delimiter(first_line: str) -> str:
    for delimiter in (";", "\t", ","):
        if delimiter in first_line:
            return delimiter
    return ";"


def detect_encoding_path(path: Path) -> str:
    """Кодировка по началу файла. Обрезанный хвост UTF-8 на границе блока не считается ошибкой."""
    with path.open("rb") as fh:
        sample = fh.read(65536)
    if not sample:
        raise ConversionError("Файл пуст")
    for encoding in ENCODINGS:
        chunk = sample
        if encoding.startswith("utf-8") and len(sample) > 3:
            chunk = sample[:-3]
        try:
            chunk.decode(encoding)
        except UnicodeDecodeError:
            continue
        return encoding
    raise ConversionError("Не удалось определить кодировку (ожидается UTF-8 или Windows-1251)")


def read_header(text_file) -> tuple[list[str], csv.reader]:
    """Заголовок и читатель строк. Файл не загружается целиком."""
    first = text_file.readline()
    if not first or not first.strip():
        raise ConversionError("Файл пуст")
    reader = csv.reader(itertools.chain([first], text_file), delimiter=detect_delimiter(first))
    try:
        headers = [cell.strip() for cell in next(reader)]
    except StopIteration as exc:
        raise ConversionError("Файл пуст") from exc
    return headers, reader


def read_rows(text: str) -> tuple[list[str], list[list[str]]]:
    first_line = text.split("\n", 1)[0]
    reader = csv.reader(io.StringIO(text), delimiter=detect_delimiter(first_line))
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    if not rows:
        raise ConversionError("Файл пуст")
    return [h.strip() for h in rows[0]], rows[1:]


@dataclass
class ColumnBinding:
    index: int
    raw_key: str
    spec: FieldSpec | None


def bind_columns(headers: list[str], entity_map: EntityMap) -> tuple[list[ColumnBinding], list[str]]:
    """Сопоставляет колонки со спецификацией: COLUMN_NAME или русский заголовок.

    Повтор одного имени берёт следующую ещё не занятую спецификацию в порядке карты.
    Так два ACCOUNT_ID остаются «код ЛС» и «признак ЧУП», а второй «Код схемы»
    в отчёте не затирает первый.
    """
    by_header = entity_map.by_header()
    used: set[int] = set()
    seen: dict[str, int] = {}
    bindings, unknown = [], []
    for index, header in enumerate(headers):
        occurrence = seen.get(header, 0)
        seen[header] = occurrence + 1
        raw_key = header if occurrence == 0 else f"{header}#{occurrence + 1}"
        spec = None
        for candidate in by_header.get(norm_header(header), []):
            if id(candidate) not in used:
                spec = candidate
                used.add(id(candidate))
                break
        if spec is None:
            unknown.append(raw_key)
        bindings.append(ColumnBinding(index, raw_key, spec))
    return bindings, unknown


def convert(value: str, kind: str):
    value = (value or "").strip()
    if value == "":
        return False if kind == "bool" else ("" if kind in {"str", "payment_type"} else None)
    try:
        if kind == "str":
            return value
        if kind in {"int", "months"}:
            compact = value.replace(" ", "")
            if _SCIENCE.match(compact):
                if kind == "int":
                    return None
                raise ConversionError(
                    f"Число «{value}» записано в экспоненциальном виде: "
                    "точные цифры потеряны при выгрузке из Excel"
                )
            return int(Decimal(compact.replace(",", ".")))
        if kind == "dec":
            return Decimal(value.replace(" ", "").replace("\u00a0", "").replace(",", "."))
        if kind == "bool":
            return value.lower() in TRUE_VALUES
        if kind == "date":
            return parse_date(value)
        if kind == "payment_type":
            mapped = PAYMENT_TYPES.get(value.lower())
            if mapped is None:
                raise ConversionError(f"Неизвестный тип оплаты «{value}»")
            return mapped
    except (InvalidOperation, ValueError) as exc:
        if isinstance(exc, ConversionError):
            raise
        raise ConversionError(f"Некорректное значение «{value}» для типа {kind}") from exc
    raise ConversionError(f"Неизвестный тип {kind}")


def excel_text(value) -> str:
    """Значение ячейки Excel в ту же строку, которую разбирает convert."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, datetime):
        if value.time() == time.min:
            return value.strftime("%d.%m.%Y")
        return value.strftime("%d.%m.%Y %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, Decimal):
        return _trim_number(value)
    if isinstance(value, float):
        if value != value or value in {float("inf"), float("-inf")}:
            return ""
        return _trim_number(Decimal(str(value)))
    return str(value).strip()


def _trim_number(value: Decimal) -> str:
    text = format(value, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text


def open_table(content: bytes) -> tuple[str, list[str], list[list[str]]]:
    """Кодировка (или xlsx/xls), заголовки и строки. CSV читается целиком, как и книга Excel."""
    if content.startswith(b"PK"):
        headers, rows = read_xlsx(content)
        return "xlsx", headers, rows
    if content.startswith(b"\xd0\xcf\x11\xe0"):
        headers, rows = read_xls(content)
        return "xls", headers, rows
    text, encoding = decode(content)
    headers, rows = read_rows(text)
    return encoding, headers, rows


def read_xlsx(content: bytes) -> tuple[list[str], list[list[str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ConversionError("Для файлов Excel нужен пакет openpyxl") from exc
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ConversionError(f"Не удалось прочитать Excel: {exc}") from exc
    try:
        if not workbook.worksheets:
            raise ConversionError("В книге Excel нет листов")
        sheet = workbook.worksheets[0]
        cursor = sheet.iter_rows(values_only=True)
        try:
            header_row = next(cursor)
        except StopIteration as exc:
            raise ConversionError("Файл пуст") from exc
        headers = [excel_text(cell).strip() for cell in header_row]
        if not any(headers):
            raise ConversionError("Файл пуст")
        rows = []
        for raw in cursor:
            row = [excel_text(cell) for cell in raw]
            if any(cell.strip() for cell in row):
                rows.append(row)
        return headers, rows
    finally:
        workbook.close()


def read_xls(content: bytes) -> tuple[list[str], list[list[str]]]:
    try:
        import xlrd
    except ImportError as exc:
        raise ConversionError("Для файлов .xls нужен пакет xlrd") from exc
    try:
        book = xlrd.open_workbook(file_contents=content)
    except xlrd.XLRDError as exc:
        raise ConversionError(f"Не удалось прочитать Excel: {exc}") from exc
    if book.nsheets == 0:
        raise ConversionError("В книге Excel нет листов")
    sheet = book.sheet_by_index(0)
    if sheet.nrows == 0:
        raise ConversionError("Файл пуст")

    def cell_text(row_index: int, column_index: int) -> str:
        cell = sheet.cell(row_index, column_index)
        if cell.ctype == xlrd.XL_CELL_DATE:
            return excel_text(xlrd.xldate_as_datetime(cell.value, book.datemode))
        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
            return "1" if cell.value else "0"
        if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
            return ""
        return excel_text(cell.value)

    headers = [cell_text(0, column).strip() for column in range(sheet.ncols)]
    if not any(headers):
        raise ConversionError("Файл пуст")
    rows = []
    for row_index in range(1, sheet.nrows):
        row = [cell_text(row_index, column) for column in range(sheet.ncols)]
        if any(cell.strip() for cell in row):
            rows.append(row)
    return headers, rows


def parse_date(value: str) -> date:
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    raise ConversionError(f"Некорректная дата «{value}»")
